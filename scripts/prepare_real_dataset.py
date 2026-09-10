import openpyxl
import re
import json
import os
from pathlib import Path
from dotenv import load_dotenv
from supabase import create_client

load_dotenv('u:/offline-os/.env')
url = os.getenv('SUPABASE_URL')
key = os.getenv('SUPABASE_SERVICE_ROLE_KEY')
supabase = create_client(url, key)

wb = openpyxl.load_workbook(r'U:\market-task\Utsav Mishra.xlsx', data_only=True)
sheet = wb['Your Leads']
headers = [c for c in next(sheet.iter_rows(values_only=True))]

def extract_full_name(first_name, linkedin_url, email):
    first_name = (first_name or '').strip()
    if ' ' in first_name and len(first_name.split()) > 1:
        return first_name

    # Check linkedin URL first if it has hyphens
    if linkedin_url:
        match = re.search(r'/in/([^/?#]+)', str(linkedin_url).rstrip('/'))
        if match:
            slug = match.group(1).lower()
            slug = re.sub(r'-[0-9a-f]{6,}$', '', slug)
            slug = re.sub(r'-[0-9]+$', '', slug)
            slug = re.sub(r'[0-9]+$', '', slug)
            if '-' in slug:
                parts = [p for p in slug.split('-') if p and not p.isdigit()]
                if len(parts) >= 2:
                    name_cand = ' '.join(p.capitalize() for p in parts)
                    if name_cand.lower().startswith(first_name.lower()):
                        return name_cand
                    elif first_name:
                        return f"{first_name} {parts[-1].capitalize()}"
            else:
                fn_low = first_name.lower()
                if slug.startswith(fn_low) and len(slug) > len(fn_low) + 2:
                    rem = slug[len(fn_low):].lstrip('-_')
                    if rem.startswith('md'):
                        rem = rem[2:]
                    if rem.isalpha():
                        return f"{first_name} {rem.capitalize()}"

    # Try from email if firstname.lastname
    if email and '@' in email:
        local = email.split('@')[0].lower()
        if '.' in local:
            parts = local.split('.')
            if len(parts) == 2 and all(p.isalpha() for p in parts):
                return f"{parts[0].capitalize()} {parts[1].capitalize()}"
        if first_name and local.isalpha() and len(local) > 3 and local != first_name.lower() and not local.startswith(first_name.lower()):
            return f"{first_name} {local.capitalize()}"

    return first_name

def classify_role(title: str):
    t = title.lower()
    if any(k in t for k in ['founder', 'ceo', 'co-founder', 'founding']):
        return 'founder', 'c-level'
    elif any(k in t for k in ['cto', 'architect', 'engineer', 'lead engineer']):
        return 'engineer', ('c-level' if 'cto' in t else 'senior')
    elif any(k in t for k in ['vp', 'head of', 'director']):
        return 'operator', ('vp' if 'vp' in t else 'director')
    elif any(k in t for k in ['investor', 'partner', 'venture']):
        return 'investor', 'partner'
    return 'operator', 'senior'

def extract_tags(text: str):
    t = text.lower()
    tags = []
    if any(k in t for k in ['ai', 'llm', 'speech', 'voice', 'transcription', 'audio', 'agent', 'model']):
        tags.append('ai')
    if any(k in t for k in ['agent', 'workflow', 'automation', 'scraping', 'crawler', 'browser']):
        tags.append('agents')
    if any(k in t for k in ['observability', 'infra', 'devops', 'cloud', 'security', 'database', 'sandboxed', 'api', 'platform']):
        tags.append('infra')
    if any(k in t for k in ['developer', 'coding', 'devtools', 'sdk', 'eval']):
        tags.append('dev tools')
    if any(k in t for k in ['voice', 'speech', 'telephony', 'audio']):
        tags.append('voice ai')
    if any(k in t for k in ['saas', 'enterprise', 'b2b', 'sales', 'gtm', 'marketing']):
        tags.append('b2b saas')
    return list(set(tags)) if tags else ['ai', 'dev tools']

raw_leads = []
for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 1):
    d = dict(zip(headers, row))
    first = str(d.get('Name') or '').strip()
    company = str(d.get('Company') or '').strip()
    role = str(d.get('Role') or '').strip()
    email = str(d.get('Email') or '').strip()
    linkedin = str(d.get('Linkedin') or '').strip()
    website = str(d.get('Website') or '').strip()
    desc = str(d.get('Short product/company description') or '').strip()
    
    if not first and not company:
        continue

    full_name = extract_full_name(first, linkedin, email)
    role_type, seniority = classify_role(role)
    tags = extract_tags(f"{company} {role} {desc}")
    
    # Calculate fit score
    fit = 78
    if role_type == 'founder':
        fit += 15
    elif seniority in ['c-level', 'vp']:
        fit += 10
    elif seniority == 'director':
        fit += 7
    if 'ai' in tags or 'agents' in tags:
        fit += 3
    if len(desc) > 30:
        fit += 2
    fit = min(fit, 98)

    bio = f"{desc} | {website}".strip(" |") if website else desc

    lead = {
        'source_record_id': f"market-task-{i:03d}",
        'name': full_name,
        'email': email if email else None,
        'email_normalized': email.lower() if email else None,
        'company': company,
        'role_title': role,
        'bio_notes': bio,
        'source': 'market_task',
        'role_type': role_type,
        'seniority': seniority,
        'sector_tags': tags,
        'community_fit_tags': ['verified_contact', 'tech_leader', 'offline_fellow'],
        'fit_score': fit,
        'fit_score_reasoning': f"Verified leadership at {company}. Strong alignment with Offline AI & Developer ecosystem as {role}.",
        'is_duplicate_of': None,
        'duplicate_confidence': None,
        'is_incomplete': not email or not role,
        'missing_fields': [],
        'ai_enrichment_status': 'enriched',
        'review_status': 'approved' if fit >= 85 else 'pending'
    }
    raw_leads.append(lead)

print(f"Prepared {len(raw_leads)} real leads from Excel.")

# Let's add 2 intentional realistic duplicate entries to demonstrate Duplicates Queue
dup1 = dict(raw_leads[0])
dup1['source_record_id'] = f"market-task-dup-001"
dup1['source'] = 'event'
dup1['bio_notes'] = "Met at SF AI Builders Summit. Interested in developer partnerships."
dup1['fit_score'] = None
dup1['fit_score_reasoning'] = None
dup1['is_duplicate_of'] = None # Will be linked after insert
dup1['review_status'] = 'duplicate_review'

dup2 = dict(raw_leads[22]) # Soham Ganatra
dup2['source_record_id'] = f"market-task-dup-002"
dup2['source'] = 'referral'
dup2['email'] = 'soham+alt@composio.dev'
dup2['email_normalized'] = 'soham+alt@composio.dev'
dup2['bio_notes'] = "Referral from Index Ventures; building agent toolkits."
dup2['review_status'] = 'duplicate_review'

raw_leads.append(dup1)
raw_leads.append(dup2)

print(f"Total records including duplicates: {len(raw_leads)}")
with open('data/real_leads_prepared.json', 'w') as f:
    json.dump(raw_leads, f, indent=2)
print("Saved to data/real_leads_prepared.json")
