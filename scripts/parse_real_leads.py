import openpyxl
import re
import json

wb = openpyxl.load_workbook(r'U:\market-task\Utsav Mishra.xlsx', data_only=True)
sheet = wb['Your Leads']
headers = [c for c in next(sheet.iter_rows(values_only=True))]

def extract_full_name(first_name, linkedin_url, email):
    first_name = (first_name or '').strip()
    if ' ' in first_name and len(first_name.split()) > 1:
        return first_name
    
    # 1. Try from email first if format like firstname.lastname@company.com
    if email and '@' in email:
        local = email.split('@')[0].lower()
        if '.' in local:
            parts = local.split('.')
            if len(parts) == 2 and all(p.isalpha() for p in parts):
                return f"{parts[0].capitalize()} {parts[1].capitalize()}"
        # If email is just lastname (e.g. marty@augmentcode.com when first_name is Chris)
        if first_name and local.isalpha() and len(local) > 3 and local != first_name.lower() and not local.startswith(first_name.lower()):
            # check if local looks like a last name
            return f"{first_name} {local.capitalize()}"

    # 2. Try from linkedin URL
    if linkedin_url:
        match = re.search(r'/in/([^/?#]+)', str(linkedin_url).rstrip('/'))
        if match:
            slug = match.group(1).lower()
            # Clean trailing hash/numeric suffixes
            slug = re.sub(r'-[0-9a-f]{6,}$', '', slug)
            slug = re.sub(r'-[0-9]+$', '', slug)
            slug = re.sub(r'[0-9]+$', '', slug)
            
            # If slug has hyphens
            if '-' in slug:
                parts = [p for p in slug.split('-') if p and not p.isdigit()]
                if len(parts) >= 2:
                    name_cand = ' '.join(p.capitalize() for p in parts)
                    if name_cand.lower().startswith(first_name.lower()):
                        return name_cand
                    elif first_name:
                        return f"{first_name} {parts[-1].capitalize()}"
                elif len(parts) == 1 and parts[0].lower() != first_name.lower():
                    return f"{first_name} {parts[0].capitalize()}"
            else:
                # Slug has no hyphens: e.g. shahramanver, willempienaar, madisonboyd
                fn_low = first_name.lower()
                if slug.startswith(fn_low) and len(slug) > len(fn_low) + 2:
                    rem = slug[len(fn_low):].lstrip('-_')
                    if rem.startswith('md'): # e.g. joaomdmoura
                        rem = rem[2:]
                    if rem.isalpha():
                        return f"{first_name} {rem.capitalize()}"
                elif slug != fn_low and slug.isalpha() and len(slug) > 3:
                    return f"{first_name} {slug.capitalize()}"

    return first_name

records = []
for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 1):
    d = dict(zip(headers, row))
    first = str(d.get('Name') or '').strip()
    company = str(d.get('Company') or '').strip()
    role = str(d.get('Role') or '').strip()
    email = str(d.get('Email') or '').strip()
    linkedin = str(d.get('Linkedin') or '').strip()
    website = str(d.get('Website') or '').strip()
    desc = str(d.get('Short product/company description') or '').strip()
    
    full_name = extract_full_name(first, linkedin, email)
    records.append({
        'name': full_name,
        'company': company,
        'role_title': role,
        'email': email,
        'linkedin': linkedin,
        'website': website,
        'bio': desc
    })

print(f"Extracted {len(records)} real records from 'Your Leads'")
for idx, r in enumerate(records[:25], 1):
    print(f"{idx:2d}. {r['name']:25s} | {r['company']:15s} | {r['role_title']:35s} | {r['email']}")
